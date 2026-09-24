# -*- coding: utf-8 -*-
"""
اختبار: إصلاح ثغرة اقتطاع الحصص، وفحص سلامة النظام، والجدول المجمّع.
يعمل على نسخة مؤقّتة — بياناتك الحقيقية لا تُمسّ.
"""
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, OSError):
    pass

import testkit  # noqa: F401
import db
import doctor
import app as webapp

FAILED = []


def ok(label, cond, extra=""):
    print("  %-56s %s%s" % (label, "✓" if cond else "!! فشل",
                            ("  " + extra) if extra else ""))
    if not cond:
        FAILED.append(label)


def main():
    db.init_db()
    webapp.app.config["TESTING"] = True
    c = webapp.app.test_client()
    conn = db.connect()
    vid = db.active_version(conn)["id"]

    # ============================== 1) لا اقتطاع صامت لعدد الحصص
    print("=== 1) نصاب المادة مقترح لا سقف ===")
    row = conn.execute(
        "SELECT a.*, sb.default_periods dp FROM assignments a "
        "JOIN subjects sb ON sb.id = a.subject_id "
        "WHERE sb.default_periods > 0 LIMIT 1").fetchone()
    tid, sub = row["teacher_id"], row["subject_id"]
    cap = row["dp"]
    big = cap + 4

    d = c.get("/api/teacher/%d/subject/%d/sections" % (tid, sub)).get_json()
    target = d["sections"][0]
    r = c.post("/api/assign/save", json={
        "teacher_id": tid, "subject_id": sub,
        "sections": [{"id": target["id"], "periods": big,
                      "label": target["label"]}]}).get_json()
    saved = conn.execute(
        "SELECT periods_per_week p FROM assignments WHERE teacher_id=? "
        "AND subject_id=? AND section_id=?", (tid, sub, target["id"])).fetchone()
    ok("القيمة الأكبر من المقترح تُحفظ كما هي", saved and saved["p"] == big,
       "أُدخل %d فحُفظ %d" % (big, saved["p"] if saved else -1))
    ok("ويُنبَّه المستخدم على الاختلاف", bool(r.get("over")),
       (r.get("message") or "")[:60])
    ok("لا تُذكر كلمة «خُفّض» في الرسالة", "خُفّض" not in (r.get("message") or ""))

    page = c.get("/assignments/teacher").get_data(as_text=True)
    ok("الواجهة لا تقصّ القيمة تلقائياً", "e.target.value = CAP" not in page)
    ok("الواجهة تُظهر تنبيه الاختلاف", "markDiff" in page)

    # ============================== 2) فحص سلامة النظام
    print("\n=== 2) فحص سلامة النظام ===")
    res = doctor.run_checks(conn, vid)
    codes = {x["code"] for x in res}
    ok("الفحص يعمل ويعيد نتائج منظّمة", isinstance(res, list))

    # اصنع تناقضاً متعمّداً: غيّر عدد الإسناد دون تغيير الجدول
    a = conn.execute(
        "SELECT a.id, a.periods_per_week, a.section_id, a.subject_id, a.teacher_id "
        "FROM assignments a JOIN schedule s ON s.section_id=a.section_id "
        "AND s.subject_id=a.subject_id AND s.teacher_id=a.teacher_id "
        "WHERE s.version_id=? LIMIT 1", (vid,)).fetchone()
    if a:
        conn.execute("UPDATE assignments SET periods_per_week=? WHERE id=?",
                     (a["periods_per_week"] + 7, a["id"]))
        conn.commit()
        res = doctor.run_checks(conn, vid)
        found = [x for x in res if x["code"] == "assign_vs_schedule"]
        ok("يكشف اختلاف الإسناد عن الجدول", len(found) == 1 and found[0]["rows"])
        done = doctor.apply_fixes(conn, vid, res)
        ok("الإصلاح التلقائي ينفّذ", bool(done), "؛ ".join(done)[:60])
        res2 = doctor.run_checks(conn, vid)
        ok("اختفى التناقض بعد الإصلاح",
           not [x for x in res2 if x["code"] == "assign_vs_schedule"])
        ok("الإسناد غير المجدول تنبيه لا خطأ",
           all(x["level"] == "warn" for x in res2
               if x["code"] == "not_scheduled_yet"))

    # خلية يتيمة
    cell = conn.execute("SELECT * FROM schedule WHERE version_id=? LIMIT 1",
                        (vid,)).fetchone()
    if cell:
        conn.execute("DELETE FROM assignments WHERE section_id=? AND subject_id=? "
                     "AND teacher_id=?",
                     (cell["section_id"], cell["subject_id"], cell["teacher_id"]))
        conn.commit()
        res = doctor.run_checks(conn, vid)
        ok("يكشف الخلايا اليتيمة",
           any(x["code"] == "orphan_cells" for x in res))
        doctor.apply_fixes(conn, vid, res)
        res = doctor.run_checks(conn, vid)
        ok("الإصلاح يعيد الإسناد المفقود",
           not any(x["code"] == "orphan_cells" for x in res))

    # إسناد مكرّر
    a = conn.execute("SELECT * FROM assignments LIMIT 1").fetchone()
    conn.execute("INSERT INTO assignments(section_id, subject_id, teacher_id, "
                 "periods_per_week) VALUES (?,?,?,?)",
                 (a["section_id"], a["subject_id"], a["teacher_id"], 1))
    conn.commit()
    res = doctor.run_checks(conn, vid)
    ok("يكشف الإسناد المكرّر", any(x["code"] == "dup_assignments" for x in res))
    doctor.apply_fixes(conn, vid, res)
    res = doctor.run_checks(conn, vid)
    ok("الإصلاح يدمج المكرّر", not any(x["code"] == "dup_assignments" for x in res))

    ok("صفحة الفحص تعمل", c.get("/doctor").status_code == 200)
    ok("زر الإصلاح يعمل", c.post("/doctor/fix").status_code in (200, 302))

    # ============================== 3) الجدول المجمّع
    print("\n=== 3) الجدول المجمّع ===")
    for kind in ("sections", "teachers"):
        r = c.get("/grid?view=%s&layout=combined" % kind)
        ok("عرض مجمّع (%s)" % kind, r.status_code == 200)
        html = r.get_data(as_text=True)
        ok("  يحوي جدولاً واحداً بكل الأيام (%s)" % kind,
           html.count("dayhead") >= 5, "%d يوم" % html.count("dayhead"))
        ok("  فيه عمود مجموع (%s)" % kind, "المجموع" in html)

        r = c.get("/grid?view=%s&layout=separate" % kind)
        ok("عرض منفصل ما زال يعمل (%s)" % kind, r.status_code == 200)
        ok("  فيه زرّا التبديل (%s)" % kind,
           "الجدول المجمّع" in r.get_data(as_text=True))

        r = c.get("/print?kind=%s&layout=combined&sig=1" % kind)
        ok("معاينة طباعة مجمّعة (%s)" % kind, r.status_code == 200)

        r = c.get("/export/excel?kind=%s&layout=combined" % kind)
        ok("Excel مجمّع (%s)" % kind, r.status_code == 200 and len(r.data) > 5000,
           "%d بايت" % len(r.data))

        r = c.get("/export/pdf?kind=%s&layout=combined&sig=1" % kind)
        ok("PDF مجمّع (%s)" % kind,
           r.status_code == 200 and r.data[:5] == b"%PDF-",
           "%d بايت" % len(r.data))

        r = c.get("/export/excel?kind=%s" % kind)
        ok("Excel المنفصل ما زال يعمل (%s)" % kind, r.status_code == 200)
        r = c.get("/export/pdf?kind=%s&sig=1" % kind)
        ok("PDF المنفصل ما زال يعمل (%s)" % kind,
           r.status_code == 200 and r.data[:5] == b"%PDF-")

    # ============================== 3ب) ألوان PDF ووضوحه
    print("\n=== 3ب) ألوان PDF ===")
    import os, tempfile, pdfplumber, exports

    def hexof(r):
        col = r.get("non_stroking_color")
        if not isinstance(col, (list, tuple)) or len(col) != 3:
            return None
        return "".join("%02X" % round(x * 255) for x in col)

    # جدول لكل فصل
    owners, cells, days, max_p, slots, stt = exports._fetch(vid, "sections")
    path = os.path.join(tempfile.gettempdir(), "t_sep.pdf")
    open(path, "wb").write(exports.build_pdf(vid, "sections").getvalue())
    o = owners[0]
    with pdfplumber.open(path) as pdf:
        pg = pdf.pages[0]
        t = pg.find_tables()[0]
        rs = [r for r in pg.rects if r.get("fill")]
        g = b = 0
        for ri, d in enumerate(days, start=1):
            for pn in range(max_p, 0, -1):
                items = cells.get((o["id"], d, pn))
                if not items:
                    continue
                bb = t.rows[ri].cells[max_p - pn]
                if not bb:
                    continue
                want = exports.PALETTE[int(items[0].get("color_index") or 0) % 12]
                hit = [hexof(r) for r in rs
                       if abs(r["x0"] - bb[0]) < 2 and abs(r["top"] - bb[1]) < 2]
                g, b = (g + 1, b) if (hit and hit[0] == want) else (g, b + 1)
        ok("ألوان المواد في PDF المنفصل صحيحة", b == 0 and g > 0,
           "صحيح %d · خاطئ %d" % (g, b))
        hdr = [x for x in t.rows[0].cells if x]
        ok("عمود اليوم في أقصى اليمين", len(hdr) == max_p + 1)
    os.remove(path)

    # الجدول المجمّع
    ow, cl, dys, slt, stt2, lim, pday, colsx, tot = exports._matrix(vid, "sections")
    brk = {x["period_number"] for x in slt if x["is_break"]}
    ncols = 2 + len(colsx)
    path = os.path.join(tempfile.gettempdir(), "t_comb.pdf")
    open(path, "wb").write(exports.build_pdf_combined(vid, "sections").getvalue())
    with pdfplumber.open(path) as pdf:
        pg = pdf.pages[0]
        t = pg.find_tables()[0]
        rs = [r for r in pg.rects if r.get("fill")]
        g = b = 0
        for ri, o2 in enumerate(ow, start=2):
            if ri >= len(t.rows):
                break
            ci = 0
            for d in dys:
                lm = lim[o2["id"]].get(d, pday[d])
                for pn in range(1, pday[d] + 1):
                    idx = ncols - 3 - ci
                    ci += 1
                    if pn > lm or pn in brk:
                        continue
                    items = cl.get((o2["id"], d, pn))
                    if not items:
                        continue
                    bb = t.rows[ri].cells[idx]
                    if not bb:
                        continue
                    want = exports.PALETTE[int(items[0].get("color_index") or 0) % 12]
                    hit = [hexof(r) for r in rs
                           if abs(r["x0"] - bb[0]) < 2 and abs(r["top"] - bb[1]) < 2]
                    g, b = (g + 1, b) if (hit and hit[0] == want) else (g, b + 1)
        ok("ألوان المواد في PDF المجمّع صحيحة", b == 0 and g > 0,
           "صحيح %d · خاطئ %d" % (g, b))

        # أسماء الأيام تظهر في رأس الجدول المجمّع (الخلية المدموجة بعد العكس)
        import import_pdf as ipx
        head = [ipx.fix(pg.crop(cc).extract_text() or "") if cc else ""
                for cc in t.rows[0].cells]
        names = [x for x in head if x][::-1]
        want = [db.DAY_NAMES[d] for d in dys]
        ok("أسماء الأيام تظهر في رأس المجمّع",
           all(w in names for w in want), "، ".join(names[:7]))
        ok("ترتيب الرأس: الاسم ثم المجموع ثم الأيام",
           names[:2] == ["الفصل", "المجموع"], str(names[:2]))
    os.remove(path)

    # ============================== 3ج) خيار بلا ألوان
    print("\n=== 3ج) خيار الألوان ===")
    from openpyxl import load_workbook
    PAL = set(exports.PALETTE)

    def pdf_hits(buf):
        pp = os.path.join(tempfile.gettempdir(), "q_col.pdf")
        open(pp, "wb").write(buf.getvalue())
        n = 0
        with pdfplumber.open(pp) as pdf:
            for r in pdf.pages[0].rects:
                cc = r.get("non_stroking_color")
                if isinstance(cc, (list, tuple)) and len(cc) == 3:
                    if "".join("%02X" % round(x * 255) for x in cc) in PAL:
                        n += 1
        os.remove(pp)
        return n

    def xl_hits(buf):
        ws = load_workbook(buf).active
        n = 0
        for rw in ws.iter_rows():
            for cc in rw:
                f = cc.fill
                if f and f.fgColor and f.fgColor.rgb and str(f.fgColor.rgb)[-6:] in PAL:
                    n += 1
        return n

    for lbl, fn, hits in (
            ("PDF منفصل", lambda cl: exports.build_pdf(vid, "sections", colored=cl), pdf_hits),
            ("PDF مجمّع", lambda cl: exports.build_pdf_combined(vid, "sections", colored=cl), pdf_hits),
            ("Excel منفصل", lambda cl: exports.build_excel(vid, "sections", colored=cl), xl_hits),
            ("Excel مجمّع", lambda cl: exports.build_excel_combined(vid, "sections", colored=cl), xl_hits)):
        on, off = hits(fn(True)), hits(fn(False))
        ok("%s: ملوّن يلوّن وبلا ألوان لا يلوّن" % lbl, on > 0 and off == 0,
           "ملوّن %d · بلا %d" % (on, off))

    for u in ("/grid?view=sections&colored=0",
              "/grid?view=sections&layout=combined&colored=0",
              "/print?kind=sections&colored=0",
              "/print?kind=sections&layout=combined&colored=0"):
        html = c.get(u).get_data(as_text=True)
        ok("وضع بلا ألوان في %s" % u.split("?")[1][:34], "nocolor" in html)
    html = c.get("/grid?view=sections").get_data(as_text=True)
    ok("الملوّن هو الافتراضي", "nocolor" not in html)
    ok("زرّا التبديل ظاهران", "بلا ألوان" in html and "ملوّن" in html)

    # ============================== 4) كل الشاشات
    print("\n=== 4) كل الشاشات ===")
    for p in ("/", "/settings", "/structure", "/subjects", "/teachers",
              "/assignments", "/assignments/teacher", "/assignments/list",
              "/assignments/balance", "/assignments/rooms", "/generate",
              "/history", "/doctor", "/print?kind=sections"):
        ok("GET %s" % p, c.get(p).status_code == 200)

    conn.close()
    print("\n" + "=" * 62)
    if FAILED:
        print("فشل %d:" % len(FAILED))
        for f in FAILED:
            print("  -", f)
        return 1
    print("كل شيء يعمل ✓")
    return 0


if __name__ == "__main__":
    sys.exit(main())
