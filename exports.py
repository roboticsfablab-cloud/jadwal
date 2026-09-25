# -*- coding: utf-8 -*-
"""
تصدير الجداول: Excel عبر openpyxl، وPDF عربي عبر reportlab
مع تشكيل الحروف (arabic_reshaper) واتجاه النص (bidi).
"""
import io
import os

import db

# ----------------------------------------------------------------- Excel

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

PALETTE = ["DBEAFE", "DCFCE7", "FEF3C7", "FCE7F3", "E0E7FF", "FFE4E6",
           "CCFBF1", "F3E8FF", "FFEDD5", "ECFCCB", "CFFAFE", "FEE2E2"]

# حدود واضحة بين كل خانة وخانة، وخط أغمق عند حدود الأيام
THIN = Side(style="thin", color="7A8496")
MED = Side(style="medium", color="3A4354")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAVY = Border(left=MED, right=MED, top=THIN, bottom=THIN)


def pick(colored, tinted, grey):
    """اللون في الوضع الملوّن، ورمادي محايد في وضع بلا ألوان."""
    return tinted if colored else grey


def edge_border(is_day_edge):
    """حدّ عادي، مع خط أغمق على يسار الخانة عند نهاية اليوم."""
    if not is_day_edge:
        return BORDER
    return Border(left=MED, right=THIN, top=THIN, bottom=THIN)


def _fetch(version_id, kind):
    """يعيد (owners, cells, days, max_p, slots, settings)."""
    conn = db.connect()
    st = db.get_settings(conn)
    days = db.working_days(st)
    slots = [dict(r) for r in conn.execute(
        "SELECT * FROM time_slots ORDER BY period_number")]
    max_p = max([int(st.get("periods_per_day") or 7)] +
                [x["period_number"] for x in slots])

    rows = [dict(r) for r in conn.execute(
        "SELECT sc.*, sb.name subject_name, sb.short_name, sb.color_index, "
        "t.name teacher_name, se.name section_name, g.name grade_name, "
        "CASE WHEN se.is_default = 1 THEN g.name ELSE g.name || ' / ' || se.name END section_label "
        "FROM schedule sc "
        "JOIN subjects sb ON sb.id = sc.subject_id "
        "JOIN teachers t  ON t.id  = sc.teacher_id "
        "JOIN sections se ON se.id = sc.section_id "
        "JOIN grades   g  ON g.id  = se.grade_id "
        "WHERE sc.version_id = ?", (version_id,))]

    if kind == "sections":
        owners = [{"id": o["id"], "title": o["section_label"]}
                  for o in conn.execute(
                      "SELECT se.id, se.name, g.name grade_name, CASE WHEN se.is_default = 1 THEN g.name ELSE g.name || ' / ' || se.name END section_label FROM sections se "
                      "JOIN grades g ON g.id = se.grade_id "
                      "ORDER BY g.sort_order, se.sort_order")]
    else:
        owners = [{"id": t["id"], "title": t["name"]}
                  for t in conn.execute(
                      "SELECT * FROM teachers ORDER BY sort_order, id")]
    conn.close()

    cells = {}
    for r in rows:
        key = (r["section_id"] if kind == "sections" else r["teacher_id"],
               r["day"], r["period_number"])
        cells.setdefault(key, []).append(r)
    return owners, cells, days, max_p, slots, st


def _cell_text(items, kind):
    """نص الخانة: المادة + (المعلم أو الشعبة)."""
    out = []
    for r in items:
        if kind == "sections":
            out.append("%s\n%s" % (r["subject_name"], r["teacher_name"]))
        else:
            out.append("%s\n%s" % (r["subject_name"], r["section_label"]))
    return "\n—\n".join(out)


def _matrix(version_id, kind):
    """أعمدة الجدول المجمّع: قائمة (يوم، حصة) بترتيب العرض + حدود كل مالك."""
    owners, cells, days, max_p, slots, st = _fetch(version_id, kind)
    conn = db.connect()
    grade_of = {r["id"]: r["grade_id"] for r in
                conn.execute("SELECT id, grade_id FROM sections")}
    conn.close()
    limits = {}
    for o in owners:
        if kind == "sections":
            limits[o["id"]] = {d: db.periods_for_grade_day(
                st, grade_of.get(o["id"]), d) for d in days}
        else:
            limits[o["id"]] = {d: db.periods_for_day(st, d) for d in days}
    per_day = {d: max([limits[o["id"]].get(d, max_p) for o in owners] or [max_p])
               for d in days}
    columns = [(d, p) for d in days for p in range(1, per_day[d] + 1)]
    totals = {}
    for o in owners:
        totals[o["id"]] = sum(
            len(cells.get((o["id"], d, p), ())) for d, p in columns)
    return owners, cells, days, slots, st, limits, per_day, columns, totals


def build_excel_combined(version_id, kind="sections", colored=True):
    """جدول واحد: صف لكل فصل (أو معلمة) وعمود لكل (يوم، حصة)."""
    owners, cells, days, slots, st, limits, per_day, columns, totals = _matrix(
        version_id, kind)
    breaks = {x["period_number"] for x in slots if x["is_break"]}

    wb = Workbook()
    ws = wb.active
    ws.title = "المجمّع"
    ws.sheet_view.rightToLeft = True

    ws.cell(row=1, column=1,
            value=st.get("school_name") or "جدول المدرسة").font = Font(size=15, bold=True)
    ws.cell(row=2, column=1,
            value="الجدول العام المجمّع — %s"
                  % ("حسب الفصول" if kind == "sections" else "حسب المعلمات")
            ).font = Font(size=10)

    # العمود 1 = الاسم، العمود 2 = المجموع، ثم أعمدة الأيام
    head1, head2, first = 4, 5, 3
    ws.cell(row=head1, column=1, value="الفصل" if kind == "sections" else "المعلمة")
    ws.cell(row=head1, column=2, value="المجموع")
    ws.merge_cells(start_row=head1, start_column=1, end_row=head2, end_column=1)
    ws.merge_cells(start_row=head1, start_column=2, end_row=head2, end_column=2)

    col = first
    day_edges = []                       # آخر عمود في كل يوم - لخطّ فاصل أغمق
    for d in days:
        n = per_day[d]
        ws.cell(row=head1, column=col, value=db.DAY_NAMES[d])
        if n > 1:
            ws.merge_cells(start_row=head1, start_column=col,
                           end_row=head1, end_column=col + n - 1)
        for k in range(n):
            c = ws.cell(row=head2, column=col + k,
                        value=("فسحة" if (k + 1) in breaks else k + 1))
            c.font = Font(bold=True, size=9)
            c.alignment = Alignment(horizontal="center")
            c.border = edge_border(col + k == col + n - 1)
        cc = ws.cell(row=head1, column=col)
        cc.font = Font(bold=True)
        cc.alignment = Alignment(horizontal="center")
        cc.fill = PatternFill("solid", fgColor=pick(colored, "E0E7FF", "E6E6E6"))
        col += n
        day_edges.append(col - 1)
    last_col = col - 1

    for cell in (ws.cell(row=head1, column=1), ws.cell(row=head1, column=2)):
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E5E7EB")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = HEAVY

    r = head2 + 1
    for o in owners:
        nm = ws.cell(row=r, column=1, value=o["title"])
        nm.font = Font(bold=True, size=9)
        nm.border = BORDER
        nm.alignment = Alignment(vertical="center", wrap_text=True)

        t = ws.cell(row=r, column=2, value=totals.get(o["id"], 0))
        t.font = Font(bold=True, size=10)
        t.border = HEAVY
        t.fill = PatternFill("solid", fgColor=pick(colored, "EEF1F6", "F0F0F0"))
        t.alignment = Alignment(horizontal="center", vertical="center")

        col = first
        for d in days:
            lim = limits[o["id"]].get(d, per_day[d])
            for p in range(1, per_day[d] + 1):
                c = ws.cell(row=r, column=col)
                c.border = edge_border(col in day_edges)
                c.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
                c.font = Font(size=8)
                if p > lim:
                    c.fill = PatternFill("solid", fgColor="E2E5EC")
                elif p in breaks:
                    c.fill = PatternFill("solid", fgColor="F3F4F6")
                else:
                    items = cells.get((o["id"], d, p))
                    if items:
                        c.value = _cell_text(items, kind)
                        if colored:
                            idx = int(items[0].get("color_index") or 0) % len(PALETTE)
                            c.fill = PatternFill("solid", fgColor=PALETTE[idx])
                col += 1
        ws.row_dimensions[r].height = 30
        r += 1

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 9
    for k in range(first, last_col + 1):
        ws.column_dimensions[get_column_letter(k)].width = 11
    ws.freeze_panes = ws.cell(row=head2 + 1, column=first)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_excel(version_id, kind="sections", colored=True):
    owners, cells, days, max_p, slots, st = _fetch(version_id, kind)
    breaks = {x["period_number"] for x in slots if x["is_break"]}

    wb = Workbook()
    ws = wb.active
    ws.title = "الفصول" if kind == "sections" else "المعلمون"
    ws.sheet_view.rightToLeft = True

    title = st.get("school_name") or "جدول المدرسة"
    ws.cell(row=1, column=1, value=title).font = Font(size=16, bold=True)
    ws.cell(row=2, column=1,
            value="الجدول العام حسب %s" % ("الفصول" if kind == "sections"
                                            else "المعلمين")).font = Font(size=11)

    r = 4
    for o in owners:
        ws.cell(row=r, column=1, value=o["title"]).font = Font(size=13, bold=True)
        r += 1
        ws.cell(row=r, column=1, value="اليوم").font = Font(bold=True)
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor="E5E7EB")
        ws.cell(row=r, column=1).border = BORDER
        for p in range(1, max_p + 1):
            c = ws.cell(row=r, column=p + 1,
                        value=("فسحة" if p in breaks else "الحصة %d" % p))
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E5E7EB")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDER
        r += 1

        for d in days:
            ws.cell(row=r, column=1, value=db.DAY_NAMES[d]).font = Font(bold=True)
            ws.cell(row=r, column=1).border = BORDER
            ws.cell(row=r, column=1).alignment = Alignment(
                horizontal="center", vertical="center")
            for p in range(1, max_p + 1):
                c = ws.cell(row=r, column=p + 1)
                c.border = BORDER
                c.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
                if p in breaks:
                    c.fill = PatternFill("solid", fgColor="F3F4F6")
                    continue
                items = cells.get((o["id"], d, p))
                if items:
                    c.value = _cell_text(items, kind)
                    if colored:
                        idx = int(items[0].get("color_index") or 0) % len(PALETTE)
                        c.fill = PatternFill("solid", fgColor=PALETTE[idx])
            ws.row_dimensions[r].height = 34
            r += 1
        r += 2

    ws.column_dimensions["A"].width = 14
    for p in range(1, max_p + 1):
        ws.column_dimensions[get_column_letter(p + 1)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ------------------------------------------------------------------- PDF

from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                Paragraph, Spacer, PageBreak, Image as RLImage)

import arabic_reshaper
from bidi.algorithm import get_display

_FONT_READY = False
FONT_NAME = "ArabicFont"
FONT_BOLD = "ArabicFont-Bold"


def _ensure_font():
    """
    خط عربي للـPDF. نبدأ بالخط المرفق مع البرنامج (Amiri برخصة OFL)
    فيعمل على أي نظام — ويندوز أو خادم Linux عند النشر. وإن غاب نلجأ
    لخطوط النظام.
    """
    global _FONT_READY
    if _FONT_READY:
        return
    here = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "fonts")
    regular = bold = None

    # 1) الخط المرفق
    for name, slot in (("Amiri-Regular.ttf", "r"), ("Amiri-Bold.ttf", "b")):
        p = os.path.join(here, name)
        if os.path.exists(p):
            if slot == "r":
                regular = p
            else:
                bold = p

    # 2) خطوط النظام (ويندوز أو Linux)
    if regular is None:
        folders = [os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"),
                   "/usr/share/fonts/truetype/dejavu",
                   "/usr/share/fonts/truetype/liberation",
                   "/usr/share/fonts"]
        for folder in folders:
            for cand in ("arial.ttf", "Tahoma.ttf", "tahoma.ttf", "segoeui.ttf",
                         "DejaVuSans.ttf", "LiberationSans-Regular.ttf"):
                p = os.path.join(folder, cand)
                if regular is None and os.path.exists(p):
                    regular = p
            for cand in ("arialbd.ttf", "tahomabd.ttf", "segoeuib.ttf",
                         "DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf"):
                p = os.path.join(folder, cand)
                if bold is None and os.path.exists(p):
                    bold = p

    if regular is None:
        raise RuntimeError("لم يُعثر على خط عربي لتصدير PDF — "
                           "تأكد من وجود static/fonts/Amiri-Regular.ttf")
    pdfmetrics.registerFont(TTFont(FONT_NAME, regular))
    pdfmetrics.registerFont(TTFont(FONT_BOLD, bold or regular))
    _FONT_READY = True


LOGO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                         "static", "img", "mark.png")


def logo_flowable(height_mm=13):
    """شعار المدرسة أعلى الصفحة، إن وُجد ملفه."""
    if not os.path.exists(LOGO_PATH):
        return None
    try:
        from PIL import Image as PILImage
        w, h = PILImage.open(LOGO_PATH).size
        hh = height_mm * mm
        img = RLImage(LOGO_PATH, width=hh * (float(w) / h), height=hh)
        img.hAlign = "CENTER"
        return img
    except Exception:
        return None


def ar(text):
    """تشكيل الحروف العربية وعكس الاتجاه ليظهر النص صحيحاً في PDF."""
    if text is None:
        return ""
    text = str(text)
    if not text:
        return ""
    try:
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def build_pdf(version_id, kind="sections", signature=False, colored=True):
    _ensure_font()
    owners, cells, days, max_p, slots, st = _fetch(version_id, kind)
    breaks = {x["period_number"] for x in slots if x["is_break"]}

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            rightMargin=10 * mm, leftMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm)

    h1 = ParagraphStyle("h1", fontName=FONT_BOLD, fontSize=16, alignment=1,
                        leading=20)
    h2 = ParagraphStyle("h2", fontName=FONT_NAME, fontSize=10, alignment=1,
                        textColor=colors.HexColor("#555555"), leading=14)
    subj_style = ParagraphStyle("subj", fontName=FONT_BOLD, fontSize=9.5,
                                alignment=1, leading=11.5,
                                textColor=colors.HexColor("#15202F"))
    who_style = ParagraphStyle("who", fontName=FONT_NAME, fontSize=7.8,
                               alignment=1, leading=9.5,
                               textColor=colors.HexColor("#414B5C"))
    cell_style = ParagraphStyle("cell", fontName=FONT_NAME, fontSize=8,
                                alignment=1, leading=10)
    head_style = ParagraphStyle("head", fontName=FONT_BOLD, fontSize=9,
                                alignment=1, leading=11)
    day_style = ParagraphStyle("dy", fontName=FONT_BOLD, fontSize=10,
                               alignment=1, leading=12)
    brk_style = ParagraphStyle("brk", fontName=FONT_NAME, fontSize=7.5,
                               alignment=1, leading=9,
                               textColor=colors.HexColor("#8A93A3"))

    story = []
    school = st.get("school_name") or "جدول المدرسة"

    for n, o in enumerate(owners):
        if n:
            story.append(PageBreak())
        mark = logo_flowable(13)
        if mark is not None:
            story.append(mark)
            story.append(Spacer(1, 1.5 * mm))
        story.append(Paragraph(ar(school), h1))
        story.append(Paragraph(
            ar("الجدول الأسبوعي — %s" % o["title"]), h2))
        story.append(Spacer(1, 5 * mm))

        # ترتيب الأعمدة بصرياً من اليسار: الحصة الأخيرة ... الحصة 1، ثم اليوم
        # (فتُقرأ من اليمين: اليوم ثم الحصة 1 فما بعدها)
        header = [Paragraph(ar("فسحة" if p in breaks else "الحصة %d" % p), head_style)
                  for p in range(max_p, 0, -1)]
        header.append(Paragraph(ar("اليوم"), day_style))
        table_data = [header]
        tints = {}                        # (صف، عمود) -> لون خلفية المادة

        for ri, d in enumerate(days, start=1):
            row = []
            for p in range(max_p, 0, -1):
                ci = max_p - p            # موقع العمود في هذا الترتيب
                if p in breaks:
                    row.append(Paragraph(ar("فسحة"), brk_style))
                    continue
                items = cells.get((o["id"], d, p))
                if not items:
                    row.append("")
                    continue
                r0 = items[0]
                who = (r0["teacher_name"] if kind == "sections"
                       else r0["section_label"])
                row.append(Paragraph(
                    '%s<br/><font size="7.8" color="#414B5C">%s</font>'
                    % (ar(r0["subject_name"]), ar(who)), subj_style))
                tints[(ri, ci)] = PALETTE[
                    int(r0.get("color_index") or 0) % len(PALETTE)]
            row.append(Paragraph(ar(db.DAY_NAMES[d]), day_style))
            table_data.append(row)

        col_w = [(doc.width - 22 * mm) / max_p] * max_p + [22 * mm]
        t = Table(table_data, colWidths=col_w,
                  rowHeights=[10 * mm] + [17 * mm] * len(days))
        style = [
            ("GRID", (0, 0), (-1, -1), 0.9, colors.HexColor("#5B6577")),
            ("BOX", (0, 0), (-1, -1), 1.8, colors.HexColor("#2B3444")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BACKGROUND", (0, 0), (-1, 0),
             colors.HexColor(pick(colored, "#DDE3EE", "#E6E6E6"))),
            ("BACKGROUND", (-1, 1), (-1, -1),
             colors.HexColor(pick(colored, "#EDF0F6", "#F0F0F0"))),
            ("LINEBELOW", (0, 0), (-1, 0), 1.8, colors.HexColor("#2B3444")),
            ("LINEBEFORE", (-1, 0), (-1, -1), 1.8, colors.HexColor("#2B3444")),
            ("TOPPADDING", (0, 0), (-1, -1), 1),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]
        # ألوان المواد (تُلغى في وضع بلا ألوان)
        if colored:
            for (ri, ci), hexcol in tints.items():
                style.append(("BACKGROUND", (ci, ri), (ci, ri),
                              colors.HexColor("#" + hexcol)))
        for ci, p in enumerate(range(max_p, 0, -1)):
            if p in breaks:
                style.append(("BACKGROUND", (ci, 1), (ci, -1),
                              colors.HexColor("#EFF1F5")))
        t.setStyle(TableStyle(style))
        story.append(t)

        if signature:
            story.append(Spacer(1, 8 * mm))
            sig = Table(
                [[Paragraph(ar("توقيع رئيسة الإشراف التعليمي"), cell_style),
                  Paragraph(ar("الاسم: %s" % (st.get("manager_name") or "")), cell_style),
                  Paragraph(ar("التوقيع: ........................"), cell_style)]],
                colWidths=[doc.width / 3.0] * 3, rowHeights=[14 * mm])
            sig.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C0CC")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            story.append(sig)

    if not owners:
        story.append(Paragraph(ar("لا توجد بيانات للتصدير"), h1))

    doc.build(story)
    buf.seek(0)
    return buf


def build_pdf_combined(version_id, kind="sections", signature=False, colored=True):
    """الجدول المجمّع في صفحة واحدة عريضة (A3 أفقي)."""
    _ensure_font()
    owners, cells, days, slots, st, limits, per_day, columns, totals = _matrix(
        version_id, kind)
    breaks = {x["period_number"] for x in slots if x["is_break"]}

    page = landscape(A3)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, rightMargin=8 * mm,
                            leftMargin=8 * mm, topMargin=8 * mm,
                            bottomMargin=8 * mm)

    h1 = ParagraphStyle("h1", fontName=FONT_BOLD, fontSize=15, alignment=1, leading=19)
    h2 = ParagraphStyle("h2", fontName=FONT_NAME, fontSize=9, alignment=1,
                        textColor=colors.HexColor("#555555"), leading=12)
    cs = ParagraphStyle("cs", fontName=FONT_BOLD, fontSize=6.6, alignment=1,
                        leading=7.6, textColor=colors.HexColor("#15202F"))
    hs = ParagraphStyle("hs", fontName=FONT_BOLD, fontSize=7.4, alignment=1, leading=9)
    ns = ParagraphStyle("ns", fontName=FONT_BOLD, fontSize=7.2, alignment=2, leading=8.6)

    story = []
    mark = logo_flowable(12)
    if mark is not None:
        story += [mark, Spacer(1, 1.2 * mm)]
    story += [Paragraph(ar(st.get("school_name") or "جدول المدرسة"), h1),
              Paragraph(ar("الجدول العام المجمّع — %s"
                           % ("حسب الفصول" if kind == "sections" else "حسب المعلمات")),
                        h2),
              Spacer(1, 4 * mm)]

    # صفّا العنوان: الأيام ثم أرقام الحصص. الأعمدة تُعكس لأن الجدول عربي.
    row_day, row_num = [], []
    for d in days:
        n = per_day[d]
        # اسم اليوم في آخر خانات مجموعته: بعد عكس الصف يصير أول خانة في
        # الدمج، وهي وحدها التي يعرض ReportLab محتواها.
        row_day.extend([""] * (n - 1))
        row_day.append(Paragraph(ar(db.DAY_NAMES[d]), hs))
        for p in range(1, n + 1):
            row_num.append(Paragraph(ar("فسحة" if p in breaks else str(p)), hs))
    head_label = Paragraph(ar("الفصل" if kind == "sections" else "المعلمة"), hs)
    total_label = Paragraph(ar("المجموع"), hs)

    # الترتيب المنطقي: الاسم، ثم المجموع، ثم أعمدة الأيام
    data = [[head_label, total_label] + row_day,
            ["", ""] + row_num]

    body_rows = []
    tints = {}                     # (صف، عمود بعد العكس) -> لون المادة
    ncols_tmp = 2 + len(columns)
    for ri, o in enumerate(owners, start=2):
        line = []
        ci = 0                     # ترتيب العمود قبل العكس (0 = أول حصة)
        for d in days:
            lim = limits[o["id"]].get(d, per_day[d])
            for p in range(1, per_day[d] + 1):
                if p > lim or p in breaks:
                    line.append("")
                    ci += 1
                    continue
                items = cells.get((o["id"], d, p))
                if items:
                    r0 = items[0]
                    who = (r0["teacher_name"] if kind == "sections"
                           else r0["section_label"])
                    line.append(Paragraph(
                        '%s<br/><font size="5.8" color="#414B5C">%s</font>'
                        % (ar(r0.get("short_name") or r0["subject_name"]), ar(who)),
                        cs))
                    # بعد عكس الصف كاملاً: العمود ci يصير ncols-1-(ci+2)
                    tints[(ri, ncols_tmp - 3 - ci)] = PALETTE[
                        int(r0.get("color_index") or 0) % len(PALETTE)]
                else:
                    line.append("")
                ci += 1
        data.append([Paragraph(ar(o["title"]), ns),
                     Paragraph(str(totals.get(o["id"], 0)), hs)] + line)
        body_rows.append(o)

    name_w = 30 * mm
    total_w = 12 * mm
    cell_w = (doc.width - name_w - total_w) / max(1, len(columns))
    widths = [name_w, total_w] + [cell_w] * len(columns)

    # العربية من اليمين: نعكس الأعمدة وعروضها معاً
    data = [list(reversed(r)) for r in data]
    widths = list(reversed(widths))

    heights = [6.5 * mm, 5.5 * mm] + [11 * mm] * len(owners)
    t = Table(data, colWidths=widths, rowHeights=heights, repeatRows=2)

    # بعد العكس: آخر عمودين هما المجموع ثم الاسم (من اليمين)
    ncols = len(widths)
    col_total = ncols - 2
    style = [
        ("GRID", (0, 0), (-1, -1), 0.8, colors.HexColor("#5B6577")),
        ("BOX", (0, 0), (-1, -1), 1.8, colors.HexColor("#2B3444")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (-1, 1),
         colors.HexColor(pick(colored, "#DDE3F5", "#E6E6E6"))),
        ("BACKGROUND", (col_total, 2), (col_total, -1),
         colors.HexColor(pick(colored, "#E7EBF3", "#F0F0F0"))),
        ("BACKGROUND", (ncols - 1, 2), (ncols - 1, -1),
         colors.HexColor(pick(colored, "#F2F4F9", "#F7F7F7"))),
        ("LINEBELOW", (0, 1), (-1, 1), 1.8, colors.HexColor("#2B3444")),
        # خط أغمق يفصل عمودي الاسم والمجموع عن أعمدة الأيام
        ("LINEBEFORE", (col_total, 0), (col_total, -1), 1.8,
         colors.HexColor("#2B3444")),
        ("SPAN", (ncols - 1, 0), (ncols - 1, 1)),      # الاسم
        ("SPAN", (col_total, 0), (col_total, 1)),      # المجموع
        ("TOPPADDING", (0, 0), (-1, -1), 0.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0.5),
        ("LEFTPADDING", (0, 2), (col_total - 1, -1), 1),
        ("RIGHTPADDING", (0, 2), (col_total - 1, -1), 1),
    ]
    # ألوان المواد - نفس ألوان الشاشة وExcel (تُلغى في وضع بلا ألوان)
    if colored:
        for (ri, ci), hexcol in tints.items():
            style.append(("BACKGROUND", (ci, ri), (ci, ri),
                          colors.HexColor("#" + hexcol)))
    # دمج أسماء الأيام + خط فاصل أغمق بين كل يوم والذي يليه
    left = 0
    for d in reversed(days):
        n = per_day[d]
        style.append(("SPAN", (left, 0), (left + n - 1, 0)))
        if left > 0:
            style.append(("LINEBEFORE", (left, 0), (left, -1), 1.4,
                          colors.HexColor("#3A4354")))
        left += n
    # تظليل الخانات خارج الدوام
    for ri, o in enumerate(body_rows, start=2):
        ci = 0
        for d in reversed(days):
            n = per_day[d]
            lim = limits[o["id"]].get(d, n)
            for p in range(n, 0, -1):
                if p > lim:
                    style.append(("BACKGROUND", (ci, ri), (ci, ri),
                                  colors.HexColor("#E2E5EC")))
                elif p in breaks:
                    style.append(("BACKGROUND", (ci, ri), (ci, ri),
                                  colors.HexColor("#F3F4F6")))
                ci += 1
    t.setStyle(TableStyle(style))
    story.append(t)

    if signature:
        story.append(Spacer(1, 6 * mm))
        sig = Table([[Paragraph(ar("رئيسة الإشراف التعليمي: %s"
                                   % (st.get("manager_name") or "")), cs),
                      Paragraph(ar("التوقيع: ................"), cs),
                      Paragraph(ar("التاريخ:     /     / 14   هـ"), cs)]],
                    colWidths=[doc.width / 3.0] * 3, rowHeights=[12 * mm])
        sig.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C0CC")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story.append(sig)

    if not owners:
        story.append(Paragraph(ar("لا توجد بيانات للتصدير"), h1))

    doc.build(story)
    buf.seek(0)
    return buf


# ------------------------------------------------------ تصدير الإسنادات
#
# نصاب المعلمات: صف لكل (معلمة، مادة، فصل)، تُدمج خانات المعلمة ونصابها
# رأسياً على كل صفوفها، وتُدمج المادة على فصولها.
# المطلوب والمُسند: صف لكل فصل مع حالته.
# البيانات تُحسب في app.py وتُمرَّر هنا جاهزة.

STATUS_FILL = {"ok": "E7F6EE", "warn": "FDF6E3", "err": "FDECEC"}
STATUS_INK = {"ok": "106B45", "warn": "8A5A12", "err": "C62828"}


def load_level(original, actual):
    """حالة النصاب الفعلي مقارنة بالأصلي: ok مطابق، warn أقل، err زائد."""
    if actual > original:
        return "err"
    if actual == original:
        return "ok"
    return "warn"


def flat_teacher_rows(teachers):
    """
    يفرد بيانات المعلمات إلى صفوف مع مدى دمج كل خانة: صف لكل
    (معلمة، مادة، فصل). المعلمة بلا إسناد تأخذ صفاً واحداً فارغاً.
    """
    out = []
    n = 0
    for t in teachers:
        subjects = t["subjects"] or [{"name": "—", "color_index": None,
                                      "rows": [{"section_label": "—",
                                                "periods": None}]}]
        t_span = sum(len(sb["rows"]) for sb in subjects)
        n += 1
        first_t = True
        for sb in subjects:
            first_s = True
            for r in sb["rows"]:
                out.append({"n": n, "teacher": t, "first_t": first_t,
                            "t_span": t_span, "subject": sb, "first_s": first_s,
                            "s_span": len(sb["rows"]),
                            "section_label": r["section_label"],
                            "periods": r["periods"],
                            "level": load_level(t["original"], t["actual"])})
                first_t = first_s = False
    return out


def _xl_head(ws, st, title):
    ws.sheet_view.rightToLeft = True
    ws.cell(row=1, column=1,
            value=st.get("school_name") or "جدول المدرسة").font = Font(size=15, bold=True)
    ws.cell(row=2, column=1, value=title).font = Font(size=11)


def _xl_header_cell(ws, row, col, value, colored):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor=pick(colored, "E0E7FF", "E6E6E6"))
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    return c


def _xl_total_row(ws, r, values, colored):
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=r, column=col, value=v)
        c.font = Font(bold=True)
        c.border = Border(left=THIN, right=THIN, top=MED, bottom=MED)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = PatternFill("solid", fgColor=pick(colored, "E0E7FF", "E6E6E6"))
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)


def _xl_save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_excel_teacher_loads(teachers, st, colored=True):
    """نصاب المعلمات وإسناداتهن في ورقة Excel واحدة."""
    wb = Workbook()
    ws = wb.active
    ws.title = "نصاب المعلمات"
    _xl_head(ws, st, "نصاب المعلمات والمواد المسندة")

    h1, h2 = 4, 5
    for col, label in ((1, "م"), (2, "اسم المعلمة"), (5, "المادة"),
                       (6, "الصف والشعبة"), (7, "عدد الحصص")):
        _xl_header_cell(ws, h1, col, label, colored)
        _xl_header_cell(ws, h2, col, None, colored)
        ws.merge_cells(start_row=h1, start_column=col, end_row=h2, end_column=col)
    _xl_header_cell(ws, h1, 3, "النصاب", colored)
    _xl_header_cell(ws, h1, 4, None, colored)
    ws.merge_cells(start_row=h1, start_column=3, end_row=h1, end_column=4)
    _xl_header_cell(ws, h2, 3, "الأصلي", colored)
    _xl_header_cell(ws, h2, 4, "الفعلي", colored)

    thick_top = Border(left=THIN, right=THIN, top=MED, bottom=THIN)
    r = h2 + 1
    for row in flat_teacher_rows(teachers):
        t, sb = row["teacher"], row["subject"]
        top = thick_top if row["first_t"] else BORDER
        for col in range(1, 8):
            c = ws.cell(row=r, column=col)
            c.border = top
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        if row["first_t"]:
            span = row["t_span"]
            lvl = row["level"]
            ws.cell(row=r, column=1, value=row["n"])
            nm = ws.cell(row=r, column=2, value=t["name"])
            nm.font = Font(bold=True)
            nm.alignment = Alignment(horizontal="right", vertical="center",
                                     wrap_text=True)
            orig = ws.cell(row=r, column=3, value=t["original"])
            orig.font = Font(bold=True)
            act = ws.cell(row=r, column=4, value=t["actual"])
            act.font = Font(bold=True,
                            color=STATUS_INK[lvl] if colored else "000000")
            if colored:
                orig.fill = PatternFill("solid", fgColor="EEF1F6")
                act.fill = PatternFill("solid", fgColor=STATUS_FILL[lvl])
            if span > 1:
                for col in (1, 2, 3, 4):
                    ws.merge_cells(start_row=r, start_column=col,
                                   end_row=r + span - 1, end_column=col)
        tint = None
        if colored and sb.get("color_index") is not None:
            tint = PatternFill("solid", fgColor=PALETTE[
                int(sb["color_index"] or 0) % len(PALETTE)])
        if row["first_s"]:
            c = ws.cell(row=r, column=5, value=sb["name"])
            c.font = Font(bold=True)
            if tint:
                c.fill = tint
            if row["s_span"] > 1:
                ws.merge_cells(start_row=r, start_column=5,
                               end_row=r + row["s_span"] - 1, end_column=5)
        c6 = ws.cell(row=r, column=6, value=row["section_label"])
        c7 = ws.cell(row=r, column=7, value=row["periods"])
        if tint:
            c6.fill = tint
            c7.fill = tint
        ws.row_dimensions[r].height = 20
        r += 1

    s_orig = sum(t["original"] for t in teachers)
    s_act = sum(t["actual"] for t in teachers)
    _xl_total_row(ws, r, ["المجموع", None, s_orig, s_act, None, None, s_act],
                  colored)

    for col, w in zip("ABCDEFG", (5, 28, 9, 9, 22, 22, 11)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=h2 + 1, column=1)
    return _xl_save(wb)


def build_excel_balance(sections, st, colored=True):
    """المطلوب والمُسند والمُدرج لكل فصل."""
    wb = Workbook()
    ws = wb.active
    ws.title = "المطلوب والمسند"
    _xl_head(ws, st, "المطلوب والمُسند لكل فصل")

    h = 4
    for col, label in enumerate(BALANCE_HEADS, start=1):
        _xl_header_cell(ws, h, col, label, colored)

    r = h + 1
    for n, s in enumerate(sections, start=1):
        vals = [n, s["label"], s["required"], s["assigned"], s["scheduled"],
                s["slots"], s["status"]]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = BORDER
            c.alignment = Alignment(horizontal="right" if col == 2 else "center",
                                    vertical="center", wrap_text=True)
        ws.cell(row=r, column=2).font = Font(bold=True)
        st_cell = ws.cell(row=r, column=7)
        st_cell.font = Font(bold=True,
                            color=STATUS_INK[s["level"]] if colored else "000000")
        if colored:
            st_cell.fill = PatternFill("solid", fgColor=STATUS_FILL[s["level"]])
            ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor="EEF1F6")
        ws.row_dimensions[r].height = 20
        r += 1

    tot = balance_totals(sections)
    _xl_total_row(ws, r, ["المجموع", None, tot["required"], tot["assigned"],
                          tot["scheduled"], tot["slots"], None], colored)

    for col, w in zip("ABCDEFG", (5, 26, 11, 11, 11, 15, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=h + 1, column=1)
    return _xl_save(wb)


BALANCE_HEADS = ["م", "الفصل", "المطلوب", "المُسند", "المُدرج",
                 "الخانات المتاحة", "الحالة"]


def balance_totals(sections):
    return {k: sum(s[k] for s in sections)
            for k in ("required", "assigned", "scheduled", "slots")}


def _pdf_doc(buf):
    return SimpleDocTemplate(buf, pagesize=A4,
                             rightMargin=10 * mm, leftMargin=10 * mm,
                             topMargin=10 * mm, bottomMargin=10 * mm)


def _pdf_head(st, title):
    h1 = ParagraphStyle("h1", fontName=FONT_BOLD, fontSize=15, alignment=1, leading=19)
    h2 = ParagraphStyle("h2", fontName=FONT_NAME, fontSize=10, alignment=1,
                        textColor=colors.HexColor("#555555"), leading=14)
    story = []
    mark = logo_flowable(13)
    if mark is not None:
        story += [mark, Spacer(1, 1.5 * mm)]
    story += [Paragraph(ar(st.get("school_name") or "جدول المدرسة"), h1),
              Paragraph(ar(title), h2), Spacer(1, 5 * mm)]
    return story


def _pdf_signature(st, width):
    cs = ParagraphStyle("sg", fontName=FONT_NAME, fontSize=9, alignment=1, leading=11)
    sig = Table([[Paragraph(ar("رئيسة الإشراف التعليمي: %s"
                               % (st.get("manager_name") or "")), cs),
                  Paragraph(ar("التوقيع: ................"), cs),
                  Paragraph(ar("التاريخ:     /     / 14   هـ"), cs)]],
                colWidths=[width / 3.0] * 3, rowHeights=[12 * mm])
    sig.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C0CC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    return [Spacer(1, 7 * mm), sig]


def _hex(h):
    return colors.HexColor("#" + h)


def _pdf_table(doc, data, widths, style, repeat):
    """يعكس الأعمدة للعربية (العمود الأول منطقياً يصير أقصى اليمين)."""
    scale = doc.width / sum(widths)
    widths = [w * scale for w in reversed(widths)]
    data = [list(reversed(r)) for r in data]
    t = Table(data, colWidths=widths, repeatRows=repeat)
    t.setStyle(TableStyle(style))
    return t


def _pdf_base_style(colored, head_rows):
    last = head_rows - 1
    return [
        ("GRID", (0, 0), (-1, -1), 0.8, colors.HexColor("#5B6577")),
        ("BOX", (0, 0), (-1, -1), 1.8, colors.HexColor("#2B3444")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (-1, last), _hex(pick(colored, "DDE3F5", "E6E6E6"))),
        ("LINEBELOW", (0, last), (-1, last), 1.8, colors.HexColor("#2B3444")),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]


def _pdf_total_style(ri, colored, L):
    return [("SPAN", (L(1), ri), (L(0), ri)),
            ("LINEABOVE", (0, ri), (-1, ri), 1.8, colors.HexColor("#2B3444")),
            ("BACKGROUND", (0, ri), (-1, ri), _hex(pick(colored, "DDE3F5", "E6E6E6")))]


def build_pdf_teacher_loads(teachers, st, signature=False, colored=True):
    _ensure_font()
    buf = io.BytesIO()
    doc = _pdf_doc(buf)
    story = _pdf_head(st, "نصاب المعلمات والمواد المسندة")

    hs = ParagraphStyle("hs", fontName=FONT_BOLD, fontSize=9.5, alignment=1, leading=12)
    bs = ParagraphStyle("bs", fontName=FONT_NAME, fontSize=9, alignment=1, leading=11)
    bb = ParagraphStyle("bb", fontName=FONT_BOLD, fontSize=9.5, alignment=1, leading=11.5)
    nm = ParagraphStyle("nm", fontName=FONT_BOLD, fontSize=9.5, alignment=2, leading=11.5)

    # الترتيب المنطقي (من اليمين): م، الاسم، الأصلي، الفعلي، المادة، الصف، الحصص
    data = [[Paragraph(ar("م"), hs), Paragraph(ar("اسم المعلمة"), hs),
             "", Paragraph(ar("النصاب"), hs),
             Paragraph(ar("المادة"), hs), Paragraph(ar("الصف والشعبة"), hs),
             Paragraph(ar("عدد الحصص"), hs)],
            ["", "", Paragraph(ar("الأصلي"), hs), Paragraph(ar("الفعلي"), hs),
             "", "", ""]]
    # موقع العمود المنطقي k بعد العكس. ReportLab يعرض محتوى أول خانة في
    # الدمج فقط، وبعد العكس تصير آخر خانة منطقية هي الأولى - فالنص فيها.
    L = lambda k: 6 - k
    style = _pdf_base_style(colored, 2)
    for k in (0, 1, 4, 5, 6):
        style.append(("SPAN", (L(k), 0), (L(k), 1)))
    style.append(("SPAN", (L(3), 0), (L(2), 0)))

    ri = 2
    for row in flat_teacher_rows(teachers):
        t, sb = row["teacher"], row["subject"]
        line = [""] * 7
        if row["first_t"]:
            lvl = row["level"]
            line[0] = Paragraph(str(row["n"]), bs)
            line[1] = Paragraph(ar(t["name"]), nm)
            line[2] = Paragraph(str(t["original"]), bb)
            line[3] = Paragraph('<font color="#%s">%d</font>' % (
                STATUS_INK[lvl] if colored else "000000", t["actual"]), bb)
            end = ri + row["t_span"] - 1
            if end > ri:
                for k in (0, 1, 2, 3):
                    style.append(("SPAN", (L(k), ri), (L(k), end)))
            if ri > 2:
                style.append(("LINEABOVE", (0, ri), (-1, ri), 1.4,
                              colors.HexColor("#2B3444")))
            if colored:
                style.append(("BACKGROUND", (L(2), ri), (L(2), end), _hex("EEF1F6")))
                style.append(("BACKGROUND", (L(3), ri), (L(3), end),
                              _hex(STATUS_FILL[lvl])))
        if row["first_s"]:
            line[4] = Paragraph(ar(sb["name"]), bb)
            if row["s_span"] > 1:
                style.append(("SPAN", (L(4), ri), (L(4), ri + row["s_span"] - 1)))
        line[5] = Paragraph(ar(row["section_label"]), bs)
        line[6] = Paragraph("" if row["periods"] is None else str(row["periods"]), bb)
        if colored and sb.get("color_index") is not None:
            style.append(("BACKGROUND", (L(6), ri), (L(4), ri),
                          _hex(PALETTE[int(sb["color_index"] or 0) % len(PALETTE)])))
        data.append(line)
        ri += 1

    s_orig = sum(t["original"] for t in teachers)
    s_act = sum(t["actual"] for t in teachers)
    data.append(["", Paragraph(ar("المجموع"), hs), Paragraph(str(s_orig), hs),
                 Paragraph(str(s_act), hs), "", "", Paragraph(str(s_act), hs)])
    style += _pdf_total_style(ri, colored, L)

    story.append(_pdf_table(doc, data, [10, 50, 17, 17, 38, 38, 20], style, 2))
    if signature:
        story += _pdf_signature(st, doc.width)
    doc.build(story)
    buf.seek(0)
    return buf


def build_pdf_balance(sections, st, signature=False, colored=True):
    _ensure_font()
    buf = io.BytesIO()
    doc = _pdf_doc(buf)
    story = _pdf_head(st, "المطلوب والمُسند لكل فصل")

    hs = ParagraphStyle("hs", fontName=FONT_BOLD, fontSize=9.5, alignment=1, leading=12)
    bs = ParagraphStyle("bs", fontName=FONT_NAME, fontSize=9.5, alignment=1, leading=11.5)
    nm = ParagraphStyle("nm", fontName=FONT_BOLD, fontSize=9.5, alignment=2, leading=11.5)

    data = [[Paragraph(ar(h), hs) for h in BALANCE_HEADS]]
    L = lambda k: 6 - k
    style = _pdf_base_style(colored, 1)
    for ri, s in enumerate(sections, start=1):
        ink = STATUS_INK[s["level"]] if colored else "000000"
        data.append([Paragraph(str(ri), bs), Paragraph(ar(s["label"]), nm),
                     Paragraph(str(s["required"]), bs),
                     Paragraph(str(s["assigned"]), bs),
                     Paragraph(str(s["scheduled"]), bs),
                     Paragraph(str(s["slots"]), bs),
                     Paragraph('<font color="#%s">%s</font>'
                               % (ink, ar(s["status"])), hs)])
        if colored:
            style.append(("BACKGROUND", (L(6), ri), (L(6), ri),
                          _hex(STATUS_FILL[s["level"]])))
            style.append(("BACKGROUND", (L(3), ri), (L(3), ri), _hex("EEF1F6")))
    ri = len(data)
    tot = balance_totals(sections)
    data.append(["", Paragraph(ar("المجموع"), hs),
                 Paragraph(str(tot["required"]), hs),
                 Paragraph(str(tot["assigned"]), hs),
                 Paragraph(str(tot["scheduled"]), hs),
                 Paragraph(str(tot["slots"]), hs), ""])
    style += _pdf_total_style(ri, colored, L)

    story.append(_pdf_table(doc, data, [10, 50, 20, 20, 20, 24, 46], style, 1))
    if signature:
        story += _pdf_signature(st, doc.width)
    doc.build(story)
    buf.seek(0)
    return buf
